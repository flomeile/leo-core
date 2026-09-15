r"""Holt ein Google Sheet ohne Connector: exportiert es über die Drive API als xlsx und schreibt
auf Wunsch eine kompakte Markdown-Tabelle daneben, die ein Modell liest.

Warum es dieses Skript gibt: Ein Connector hängt an der Sitzung der App. Ein kopfloser Lauf aus
einer geplanten Aufgabe hat keine Sitzung und ist damit blind für alles, was nur über einen
Connector erreichbar ist (AGENTS.md, Abschnitt 16). Deshalb derselbe Weg wie bei jeder anderen
Live-Quelle: deterministisch abholen, als Datei ablegen, das Modell liest die Datei.

Zugang: OAuth mit dem eigenen Google-Konto des Besitzers (kein Dienstkonto), weil ein Dienstkonto
auf jedes geteilte Sheet erst freigegeben werden müsste; mit der eigenen Anmeldung sieht das
Skript genau das, was der Besitzer sieht. Anmeldung, Token und Scopes: google_zugang.py
(gemeinsam mit gmail-senden.py). Einrichtung: 10_System\Google-Anbindung.md.

Aufrufe:
  python google-sheets-fetch.py --auth
      erstmalige Anmeldung (öffnet den Browser)
  python google-sheets-fetch.py --sheet <ID> --xlsx <ziel.xlsx> [--markdown <ziel.md> [--blatt <Name>] [--zeilen 10]]
      normaler Lauf; die ID steht in der URL des Sheets (docs.google.com/spreadsheets/d/<ID>/edit).
      --markdown schreibt zusätzlich die letzten --zeilen Zeilen des Blattes als Tabelle (braucht openpyxl);
      scheitert der Lauf, steht der Fehler IN der Markdown-Datei, damit eine Lücke nicht wie eine
      unauffällige Woche aussieht.
"""

import sys
import argparse
import datetime
import pathlib

HIER = pathlib.Path(__file__).resolve()
if str(HIER.parent) not in sys.path:
    sys.path.insert(0, str(HIER.parent))
import google_zugang

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ANMELDEN = "python \"%s\" --auth" % HIER


def fehler(text, markdown=None, code=2):
    """Meldet den Fehler und schreibt ihn in die Markdown-Zieldatei, falls eine verlangt war."""
    if markdown:
        markdown.parent.mkdir(parents=True, exist_ok=True)
        markdown.write_text("# Google Sheet\n\n**Nicht abrufbar.** " + text + "\n\n"
                            "Einrichtung: `10_System\\Google-Anbindung.md`.\n", encoding="utf-8")
    print("FEHLER " + text)
    sys.exit(code)


def lade_sheet(creds, sheet_id, ziel):
    dienst = google_zugang.dienst("drive", "v3", creds)
    daten = dienst.files().export(fileId=sheet_id, mimeType=XLSX_MIME).execute()
    ziel.parent.mkdir(parents=True, exist_ok=True)
    ziel.write_bytes(daten)
    return ziel


def formatiere(wert):
    if wert is None:
        return ""
    if isinstance(wert, bool):
        return "ja" if wert else "nein"
    if isinstance(wert, (int, float)):
        return "{:,.2f}".format(wert).rstrip("0").rstrip(".").replace(",", "'")
    if isinstance(wert, (datetime.datetime, datetime.date)):
        return wert.strftime("%d.%m.%Y")
    return str(wert).replace("|", "\\|").replace("\n", " ")


def baue_markdown(xlsx, markdown, blatt=None, zeilen=10):
    """Die letzten `zeilen` belegten Zeilen des Blattes als Markdown-Tabelle, Kopfzeile aus der
    ersten Zeile. Bewusst nur Abschrift, keine Deutung: die Einordnung macht das Modell."""
    try:
        import openpyxl
    except ImportError:
        fehler("Bibliothek openpyxl fehlt (nur fuer --markdown noetig). Einmalig: pip install openpyxl", markdown)
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb[blatt] if blatt else wb.worksheets[0]
    belegte = [r for r in range(1, ws.max_row + 1) if any(ws.cell(r, c).value is not None for c in range(1, ws.max_column + 1))]
    if not belegte:
        fehler("Das Blatt %r ist leer." % ws.title, markdown)
    kopf = [formatiere(ws.cell(belegte[0], c).value) or ("Spalte %d" % c) for c in range(1, ws.max_column + 1)]
    daten = [r for r in belegte[1:]][-zeilen:]
    text = ["# Google Sheet: %s, Blatt %s" % (xlsx.stem, ws.title), "",
            "Abgerufen am %s ohne Connector, direkt aus dem Sheet (Rohdatei: `%s`). Die letzten %d von %d belegten Zeilen."
            % (datetime.datetime.now().strftime("%d.%m.%Y %H:%M"), xlsx.name, len(daten), max(len(belegte) - 1, 0)), "",
            "| " + " | ".join(kopf) + " |", "|" + "|".join(["---"] * len(kopf)) + "|"]
    for r in daten:
        text.append("| " + " | ".join(formatiere(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)) + " |")
    text.append("")
    markdown.parent.mkdir(parents=True, exist_ok=True)
    markdown.write_text("\n".join(text), encoding="utf-8")
    print("OK Markdown geschrieben: %s (%d Zeilen)" % (markdown, len(daten)))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--auth", action="store_true", help="erstmalige Anmeldung (oeffnet den Browser)")
    ap.add_argument("--sheet", help="ID des Sheets (aus der URL)")
    ap.add_argument("--xlsx", help="Zielpfad der xlsx-Datei")
    ap.add_argument("--markdown", help="optional: Zielpfad einer Markdown-Tabelle")
    ap.add_argument("--blatt", help="Blattname fuer --markdown (Standard: erstes Blatt)")
    ap.add_argument("--zeilen", type=int, default=10, help="Anzahl letzter Zeilen fuer --markdown (Standard 10)")
    a = ap.parse_args()
    markdown = pathlib.Path(a.markdown) if a.markdown else None
    try:
        creds = google_zugang.hole_credentials(interaktiv=a.auth, anmelde_hinweis=ANMELDEN)
    except google_zugang.KeineAnmeldung as e:
        fehler(str(e), markdown)
    if a.auth and not a.sheet:
        print("OK Google-Anmeldung vorhanden (%s)" % google_zugang.TOKEN_FILE)
        return
    if not (a.sheet and a.xlsx):
        fehler("--sheet <ID> und --xlsx <ziel> sind Pflicht (oder --auth).", markdown)
    try:
        xlsx = lade_sheet(creds, a.sheet, pathlib.Path(a.xlsx))
    except Exception as e:
        fehler("Export fehlgeschlagen (%s: %s). Ist die Drive API im Cloud-Projekt aktiviert und das Sheet mit dem Konto geteilt?"
               % (type(e).__name__, e), markdown)
    print("OK xlsx geschrieben: %s" % xlsx)
    if markdown:
        baue_markdown(xlsx, markdown, a.blatt, a.zeilen)


if __name__ == "__main__":
    main()
